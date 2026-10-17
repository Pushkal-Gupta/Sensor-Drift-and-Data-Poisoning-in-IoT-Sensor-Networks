"""Minimal manuscript builder on python-docx.

Text may contain citation markers {cite:key1,key2}.  They are resolved in
document order to Elsevier numbered style ([1], [2,3], [4-6]) and the reference
list is emitted in that order, so numbering always matches first appearance.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CITE = re.compile(r"\{cite:([^}]+)\}")
FMT = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|_\{[^}]+\}|\^\{[^}]+\})")


class Manuscript:
    def __init__(self, refs: Dict[str, str]):
        self.doc = Document()
        self.refs = refs
        self.order: List[str] = []
        self.n_fig = 0
        self.n_tab = 0
        self.n_eq = 0
        self._style()

    # ------------------------------------------------------------------ style
    def _style(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        for m in ("left_margin", "right_margin"):
            setattr(sec, m, Cm(2.5))
        sec.top_margin = sec.bottom_margin = Cm(2.5)
        st = d.styles["Normal"]
        st.font.name = "Times New Roman"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.5
        for lvl, size in ((1, 13), (2, 12), (3, 11)):
            h = d.styles[f"Heading {lvl}"]
            h.font.name = "Times New Roman"
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = lvl == 3
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------- citations
    def _resolve(self, text: str) -> str:
        def rep(m):
            keys = [k.strip() for k in m.group(1).split(",")]
            nums = []
            for k in keys:
                if k not in self.refs:
                    raise KeyError(f"unknown reference key {k!r}")
                if k not in self.order:
                    self.order.append(k)
                nums.append(self.order.index(k) + 1)
            return "[" + _compress(sorted(set(nums))) + "]"
        return CITE.sub(rep, text)

    # ------------------------------------------------------------------- runs
    def _runs(self, p, text: str, size: Optional[float] = None, italic=False):
        text = self._resolve(text)
        for part in FMT.split(text):
            if not part:
                continue
            if part.startswith("**"):
                r = p.add_run(part[2:-2]); r.bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                r = p.add_run(part[1:-1]); r.italic = True
            elif part.startswith("_{"):
                r = p.add_run(part[2:-1]); r.font.subscript = True
            elif part.startswith("^{"):
                r = p.add_run(part[2:-1]); r.font.superscript = True
            else:
                r = p.add_run(part)
            if italic:
                r.italic = True
            if size:
                r.font.size = Pt(size)
        return p

    # ----------------------------------------------------------------- blocks
    def title(self, text):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.bold = True; r.font.size = Pt(16)

    def centered(self, text, size=11, italic=False):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._runs(p, text, size, italic)

    def h1(self, t): self.doc.add_heading(t, level=1)
    def h2(self, t): self.doc.add_heading(t, level=2)
    def h3(self, t): self.doc.add_heading(t, level=3)

    def p(self, text, lead: Optional[str] = None):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if lead:
            r = para.add_run(lead + " "); r.bold = True
        self._runs(para, text)
        return para

    def bullets(self, items: Sequence[str], numbered=False):
        for it in items:
            para = self.doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self._runs(para, it)

    def eq(self, text: str, label: Optional[str] = None) -> int:
        self.n_eq += 1
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = t.rows[0].cells
        c0.width, c1.width = Cm(14.0), Cm(2.0)
        p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._runs(p0, text, italic=False)
        p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.add_run(f"({self.n_eq})")
        return self.n_eq

    def figure(self, path: Path, caption: str, width_cm: float = 15.5) -> int:
        self.n_fig += 1
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        c = self.doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = c.add_run(f"Fig. {self.n_fig}. "); r.bold = True; r.font.size = Pt(10)
        self._runs(c, caption, size=10)
        return self.n_fig

    def table(self, caption: str, header: Sequence[str], rows: Sequence[Sequence[str]],
              note: Optional[str] = None, widths: Optional[Sequence[float]] = None,
              font: float = 9) -> int:
        self.n_tab += 1
        c = self.doc.add_paragraph()
        r = c.add_run(f"Table {self.n_tab}. "); r.bold = True; r.font.size = Pt(10)
        self._runs(c, caption, size=10)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, h in enumerate(header):
            cell = t.rows[0].cells[j]
            cell.paragraphs[0].text = ""
            rr = cell.paragraphs[0].add_run(self._resolve(h)); rr.bold = True
            rr.font.size = Pt(font)
            _shade(cell, "E8E8E8")
        for i, row in enumerate(rows, start=1):
            for j, v in enumerate(row):
                para = t.rows[i].cells[j].paragraphs[0]
                para.paragraph_format.line_spacing = 1.0
                para.paragraph_format.space_after = Pt(0)
                self._runs(para, str(v), size=font)
        if widths:
            for row in t.rows:
                for j, w in enumerate(widths):
                    row.cells[j].width = Cm(w)
        for row in t.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.line_spacing = 1.0
                    para.paragraph_format.space_after = Pt(0)
        if note:
            n = self.doc.add_paragraph()
            self._runs(n, note, size=9, italic=False)
        else:
            self.doc.add_paragraph()
        return self.n_tab

    def page_break(self):
        self.doc.add_page_break()

    def references(self):
        self.h1("References")
        for i, k in enumerate(self.order, start=1):
            para = self.doc.add_paragraph()
            para.paragraph_format.line_spacing = 1.15
            para.paragraph_format.left_indent = Cm(0.8)
            para.paragraph_format.first_line_indent = Cm(-0.8)
            r = para.add_run(f"[{i}] "); r.font.size = Pt(10)
            r2 = para.add_run(self.refs[k]); r2.font.size = Pt(10)
        unused = [k for k in self.refs if k not in self.order]
        return unused

    def save(self, path: Path):
        self.doc.save(str(path))


def _compress(nums: List[int]) -> str:
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(f"{nums[i]}–{nums[j]}" if j - i >= 2 else
                   ",".join(str(n) for n in nums[i:j + 1]))
        i = j + 1
    return ",".join(out)


def _shade(cell, hex_fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)
